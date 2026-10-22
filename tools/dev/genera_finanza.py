"""Builds docs/finanza/mercato-e-finanza.xlsx: rules, examples, pity and charts of the stock-market design (#268).

    python tools/dev/genera_finanza.py . docs/finanza/mercato-e-finanza.xlsx   # then recalc the formulas in Excel/LibreOffice

import json
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

REPO = sys.argv[1]
OUT = sys.argv[2]

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F2A36")
YELLOW = PatternFill("solid", fgColor="FFFF00")
NOTE = Font(name=FONT, italic=True, color="52514E")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="C8C8C8")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()


def sheet(name, widths):
    ws = wb.create_sheet(name)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def title(ws, text, subtitle=None):
    ws["A1"] = text
    ws["A1"].font = TITLE
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = NOTE
        ws["A2"].alignment = Alignment(wrap_text=False)


def header(ws, row, labels):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = HEAD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX


def cell(ws, ref, value, font=BLACK, fmt=None, fill=None, comment=None):
    c = ws[ref]
    c.value = value
    c.font = font
    c.border = BOX
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, "Claude")
    return c


# ---------------------------------------------------------------- Guida
ws = wb.active
ws.title = "Guida"
for col, w in zip("ABC", (28, 70, 22)):
    ws.column_dimensions[col].width = w
ws.sheet_view.showGridLines = False
title(ws, "Mercato come una borsa e ramo Finanza — meccaniche, pity, esempi",
      "Documento di design per l'issue #268 di tIndustry (22 ottobre 2026). Non è ancora la specifica: serve a vedere i numeri insieme.")
rows = [
    ("Come si legge", ""),
    ("Testo blu", "Parametro modificabile: cambialo e tutti i fogli si ricalcolano."),
    ("Sfondo giallo", "Proposta da approvare (numero scelto da Claude, non ancora deciso dall'utente)."),
    ("Testo verde", "Valore preso da un altro foglio (di solito «Parametri»)."),
    ("Testo nero", "Formula."),
    ("", ""),
    ("Fogli", ""),
    ("Parametri", "Tutte le regole numeriche in un posto solo, con chi le ha decise."),
    ("Valori B", "Valore base (= soglia minima) per oggetto, prezzo medio, picco, jackpot per unità."),
    ("Fasi del prezzo", "Le fasi del prezzo e la curva di un picco secondo per secondo (grafico)."),
    ("Probabilità jackpot", "Probabilità di jackpot in base ai pezzi venduti (grafico)."),
    ("Pity", "Regola del pity e cinque scenari passo per passo."),
    ("Esempi di vendita", "Sette vendite calcolate: incasso, crollo, LED dello schermo, jackpot."),
    ("Simulazione 1 ora", "Un'ora di prezzo di un oggetto: andamento normale, un picco, un errore e il crollo (grafico)."),
    ("Ramo Finanza", "Schermo di borsa, bot della finanza e moduli nella Ricerca."),
    ("Magazzino CORE", "Anteprima della specifica 3: magazzino per oggetto e drone per livello del CORE."),
    ("Punti aperti", "Cose da decidere, compresa un'incoerenza trovata facendo i conti."),
    ("", ""),
    ("Le tre specifiche", ""),
    ("1. Mercato", "Prezzi, picchi, crolli, vendite, vendita automatica a B, valori B, salvataggi, schermata panoramica. Il jackpot è già calcolato nella simulazione ma scatta solo se in partita c'è uno schermo di borsa."),
    ("2. Schermo di borsa", "Blocco accanto al CORE, bot della finanza (drone 2.5D), sequenza del «Vendi tutto», moduli, ramo Finanza della Ricerca."),
    ("3. Magazzino del CORE", "Massimo per oggetto che cresce col livello, linee che si fermano a CORE pieno, drone più veloce per livello."),
]
r = 4
for a, b in rows:
    ws.cell(row=r, column=1, value=a).font = BOLD if b == "" or a in ("Testo blu", "Sfondo giallo", "Testo verde", "Testo nero") else BOLD
    ws.cell(row=r, column=2, value=b).font = BLACK
    ws.cell(row=r, column=2).alignment = WRAP
    r += 1
ws["A5"].font = Font(name=FONT, bold=True, color="0000FF")
ws["A6"].fill = YELLOW
ws["A7"].font = Font(name=FONT, bold=True, color="008000")

# ---------------------------------------------------------------- Parametri
P = sheet("Parametri", (46, 14, 12, 22, 60))
title(P, "Parametri", "Ogni regola in una cella: gli altri fogli li leggono da qui.")
header(P, 4, ["Regola", "Valore", "Unità", "Decisa da", "Note"])
params = [
    # key, label, value, unit, who, note, fmt
    ("norm_min", "Prezzo normale: minimo", 1.0, "× B", "utente", "Il valore base B è anche la soglia minima: fuori da un crollo il prezzo non va sotto.", "0.00"),
    ("norm_max", "Prezzo normale: massimo", 1.6, "× B", "proposta", "Fascia approvata dall'utente.", "0.00"),
    ("norm_avg", "Prezzo normale: media", 1.3, "× B", "proposta", "Usata per stimare i valori B dal listino di oggi.", "0.00"),
    ("spike_max", "Picco: massimo", 3.0, "× B", "utente", "+200%: ogni picco ha un massimo casuale fra la soglia del picco e questo.", "0.00"),
    ("spike_in", "Picco: soglia «nel picco»", 2.8, "× B", "utente", "+180%: da qui una vendita conta come fatta nel picco.", "0.00"),
    ("spike_up", "Picco: salita", 10, "s", "utente", "", "0"),
    ("spike_hold", "Picco: resta in cima", 10, "s", "utente", "", "0"),
    ("spike_down", "Picco: discesa", 30, "s", "utente", "", "0"),
    ("spike_every_min", "Picco: intervallo minimo", 10, "min", "utente", "Per oggetto; il conto alla rovescia resta fermo durante un crollo.", "0"),
    ("spike_every_max", "Picco: intervallo massimo", 20, "min", "utente", "", "0"),
    ("impact", "Vendita: calo per pezzo venduto", 0.001, "× B", "proposta", "0,1% di B per pezzo: 100 pezzi = −10% di B.", "0.0%"),
    ("recover", "Vendita: tempo per tornare all'andamento", 120, "s", "proposta", "", "0"),
    ("crash_qty", "Crollo: pezzi in una vendita", 400, "pezzi", "utente", "Una vendita oltre questa quantità fa crollare; crolla anche se le vendite piccole portano il prezzo a B.", "#,##0"),
    ("crash_len", "Crollo: durata", 20, "min", "utente", "Niente picchi e conto alla rovescia del picco fermo.", "0"),
    ("crash_low_len", "Crollo: minuti in cui può andare sotto B", 5, "min", "utente", "I primi minuti del crollo.", "0"),
    ("crash_low_min", "Crollo: minimo sotto B", 0.6, "× B", "utente", "30–40% sotto la soglia minima.", "0.00"),
    ("crash_low_max", "Crollo: massimo sotto B", 0.7, "× B", "utente", "", "0.00"),
    ("jp_min_qty", "Jackpot: pezzi minimi", 400, "pezzi", "utente", "Solo con «Vendi tutto», solo nel picco, solo con uno schermo di borsa.", "#,##0"),
    ("jp_base", "Jackpot: probabilità base", 0.02, "", "utente", "Alla quantità minima.", "0%"),
    ("jp_step_qty", "Jackpot: pezzi per ogni aumento", 200, "pezzi", "utente", "", "#,##0"),
    ("jp_step", "Jackpot: aumento", 0.01, "", "utente", "", "0%"),
    ("jp_cap", "Jackpot: tetto senza pity", 0.5, "", "utente", "Si raggiunge a 10.000 pezzi.", "0%"),
    ("jp_bonus", "Jackpot: bonus sul valore base", 7.77, "× B", "utente", "+777%, arrotondato per difetto per unità.", "0.00"),
    ("pity_add", "Pity: aumento per tentativo a vuoto", 0.25, "", "utente", "Si somma alla probabilità del tentativo successivo, anche oltre il tetto.", "0%"),
    ("pity_qty", "Pity: pezzi minimi perché il tentativo conti", 3000, "pezzi", "utente", "", "#,##0"),
    ("pity_guarantee", "Pity: tentativi a vuoto di fila per la garanzia", 3, "tentativi", "utente", "Il tentativo dopo è un jackpot sicuro.", "0"),
    ("auto_price", "Vendita automatica: prezzo", 1.0, "× B", "utente", "Paga B e non tocca il mercato.", "0.00"),
    ("led_mild", "LED blu: calo fino a", 0.10, "di B", "proposta", "Sotto questa soglia il calo è lieve (blu); oltre, grave (viola); crollo = giallo.", "0%"),
    ("late_mult", "Oggetti di fine partita: moltiplicatore di B", 1.5, "×", "proposta", "L'utente vuole oggetti avanzati più cari; il numero è da decidere.", "0.00"),
]
NAMES = {}
for i, (key, label, value, unit, who, note, fmt) in enumerate(params):
    row = 5 + i
    cell(P, f"A{row}", label)
    cell(P, f"B{row}", value, BLUE, fmt, YELLOW if who == "proposta" else None)
    cell(P, f"C{row}", unit)
    cell(P, f"D{row}", "utente (22/10)" if who == "utente" else "proposta da approvare", fill=YELLOW if who == "proposta" else None)
    cell(P, f"E{row}", note).alignment = WRAP
    NAMES[key] = f"Parametri!$B${row}"


def p(key):
    return NAMES[key]


# ---------------------------------------------------------------- Valori B
with open(os.path.join(REPO, "data", "content.json"), encoding="utf-8") as fh:
    content = json.load(fh)
market = content["market"]
LATE = {"graphite", "silicon", "titanium-ore", "titanium-plate"}
V = sheet("Valori B", (20, 16, 12, 12, 12, 12, 12, 12, 12, 16))
title(V, "Valori base B per oggetto",
      "Proposta: B = listino di oggi ÷ media normale (1,3), × moltiplicatore per gli oggetti di fine partita. Listino e minimo di oggi da data/content.json.")
header(V, 4, ["Oggetto", "id", "Listino oggi ($)", "Minimo oggi ($)", "Fine partita", "B proposto ($)",
              "Prezzo medio ($)", "Picco max ($)", "Soglia picco ($)", "Jackpot al picco max ($/pezzo)"])
for i, m in enumerate(market):
    r = 5 + i
    cell(V, f"A{r}", m["displayName"])
    cell(V, f"B{r}", m["itemId"])
    cell(V, f"C{r}", m["sellPrice"], fmt="$#,##0", comment="Fonte: data/content.json, market." if i == 0 else None)
    cell(V, f"D{r}", m.get("minSellPrice"), fmt="$#,##0")
    cell(V, f"E{r}", "sì" if m["itemId"] in LATE else "no", BLUE, fill=YELLOW)
    cell(V, f"F{r}", f'=MAX(1,ROUND(C{r}/{p("norm_avg")}*IF(E{r}="sì",{p("late_mult")},1),0))', fmt="$#,##0", fill=YELLOW)
    cell(V, f"G{r}", f"=ROUND(F{r}*{p('norm_avg')},0)", fmt="$#,##0")
    cell(V, f"H{r}", f"=ROUND(F{r}*{p('spike_max')},0)", fmt="$#,##0")
    cell(V, f"I{r}", f"=ROUND(F{r}*{p('spike_in')},0)", fmt="$#,##0")
    cell(V, f"J{r}", f"=F{r}+ROUNDDOWN(F{r}*{p('jp_bonus')},0)+(H{r}-F{r})", fmt="$#,##0")
last_b = 4 + len(market)
note_r = last_b + 2
V[f"A{note_r}"] = "Il jackpot per pezzo vale B + 777% di B (per difetto) + (prezzo del momento − B). Esempio dell'utente: B 10, prezzo 30 → 10 + 77 + 20 = 107."
V[f"A{note_r}"].font = NOTE
ch = BarChart()
ch.type = "bar"
ch.style = 10
ch.title = "Valore base B e picco massimo per oggetto"
ch.y_axis.title = "$"
ch.add_data(Reference(V, min_col=6, min_row=4, max_row=last_b), titles_from_data=True)
ch.add_data(Reference(V, min_col=8, min_row=4, max_row=last_b), titles_from_data=True)
ch.set_categories(Reference(V, min_col=1, min_row=5, max_row=last_b))
ch.series[0].graphicalProperties.solidFill = "2A78D6"
ch.series[1].graphicalProperties.solidFill = "EB6834"
ch.height, ch.width = 10, 20
V.add_chart(ch, f"A{note_r + 2}")

# ---------------------------------------------------------------- Fasi del prezzo
F = sheet("Fasi del prezzo", (26, 18, 18, 50))
title(F, "Fasi del prezzo (tutto in multipli di B)", "Ogni oggetto ha il suo mercato: le fasi di un oggetto non toccano gli altri.")
header(F, 4, ["Fase", "Da (× B)", "A (× B)", "Cosa succede"])
phases = [
    ("Normale", f"={p('norm_min')}", f"={p('norm_max')}", "Oscilla dolcemente, cambi visibili ogni pochi secondi. Ogni pezzo venduto abbassa il prezzo dello 0,1% di B, poi in circa 2 minuti torna al suo andamento."),
    ("Picco: salita", "=prezzo del momento", f"={p('spike_max')}", "In 10 s verso un massimo casuale fra 2,8×B e 3×B."),
    ("Picco: nel picco", f"={p('spike_in')}", f"={p('spike_max')}", "Finestra per vendere: da +180% la vendita conta come fatta nel picco (jackpot possibile)."),
    ("Picco: discesa", f"={p('spike_max')}", "=andamento normale", "In 30 s torna nella fascia normale."),
    ("Crollo: primi 5 minuti", f"={p('crash_low_min')}", f"={p('norm_min')}", "Vicino a B, ogni tanto 30–40% sotto."),
    ("Crollo: resto (fino a 20 min)", f"={p('norm_min')}", f"={p('norm_min')}+0.05", "Fermo vicino a B. Niente picchi, conto alla rovescia del picco fermo."),
]
for i, (a, b, c, d) in enumerate(phases):
    r = 5 + i
    cell(F, f"A{r}", a)
    cell(F, f"B{r}", b if b.startswith("=") and "prezzo" not in b and "andamento" not in b else b.lstrip("="),
         GREEN if b.startswith("=") and "prezzo" not in b and "andamento" not in b else BLACK, "0.00")
    cell(F, f"C{r}", c if c.startswith("=") and "andamento" not in c else c.lstrip("="),
         GREEN if c.startswith("=") and "andamento" not in c else BLACK, "0.00")
    cell(F, f"D{r}", d).alignment = WRAP

r0 = 13
F[f"A{r0 - 1}"] = "Un picco secondo per secondo (parte da 1,3×B, arriva a 3×B)"
F[f"A{r0 - 1}"].font = BOLD
header(F, r0, ["Secondo", "Prezzo (× B)", "Nel picco?", ""])
for t in range(0, 61):
    r = r0 + 1 + t
    cell(F, f"A{r}", t)
    up, hold, down = p("spike_up"), p("spike_hold"), p("spike_down")
    start, top = p("norm_avg"), p("spike_max")
    formula = (f"=IF(A{r}<={up},{start}+({top}-{start})*A{r}/{up},"
               f"IF(A{r}<={up}+{hold},{top},"
               f"IF(A{r}<={up}+{hold}+{down},{top}-({top}-{start})*(A{r}-{up}-{hold})/{down},{start})))")
    cell(F, f"B{r}", formula, fmt="0.00")
    cell(F, f"C{r}", f'=IF(B{r}>={p("spike_in")},"sì","")')
last_t = r0 + 61
cell(F, f"D{r0 + 1}", f'=COUNTIF(C{r0 + 1}:C{last_t},"sì")&" secondi nel picco (finestra per vendere)"', BOLD)
lc = LineChart()
lc.title = "Un picco: salita 10 s, in cima 10 s, discesa 30 s"
lc.y_axis.title = "Prezzo (× B)"
lc.x_axis.title = "Secondi"
lc.add_data(Reference(F, min_col=2, min_row=r0, max_row=last_t), titles_from_data=True)
lc.set_categories(Reference(F, min_col=1, min_row=r0 + 1, max_row=last_t))
lc.series[0].graphicalProperties.line.solidFill = "2A78D6"
lc.series[0].graphicalProperties.line.width = 25000
lc.series[0].smooth = False
lc.y_axis.scaling.min = 0
lc.y_axis.scaling.max = 3.2
lc.legend = None
lc.height, lc.width = 9, 18
F.add_chart(lc, "E4")

# ---------------------------------------------------------------- Probabilità jackpot
J = sheet("Probabilità jackpot", (16, 20, 50))
title(J, "Probabilità di jackpot con «Vendi tutto» nel picco",
      "2% a 400 pezzi, +1% ogni 200 pezzi in più, tetto 50% a 10.000. Sotto 400 pezzi: nessun jackpot. Il pity si aggiunge sopra (foglio Pity).")
header(J, 4, ["Pezzi venduti", "Probabilità", "Note"])
qty = [0, 200, 399, 400, 600, 1000, 2000, 3000, 4000, 5000, 6000, 7000, 8000, 9000, 9800, 10000, 12000]
for i, q in enumerate(qty):
    r = 5 + i
    cell(J, f"A{r}", q, BLUE, "#,##0")
    cell(J, f"B{r}",
         f"=IF(A{r}<{p('jp_min_qty')},0,MIN({p('jp_cap')},{p('jp_base')}+INT((A{r}-{p('jp_min_qty')})/{p('jp_step_qty')})*{p('jp_step')}))",
         fmt="0%")
    cell(J, f"C{r}", f'=IF(A{r}<{p("jp_min_qty")},"sotto la soglia",IF(A{r}>={p("pity_qty")},"conta anche per il pity",""))')
last_q = 4 + len(qty)
# Fine-grained series for the chart.
cr0 = last_q + 3
J[f"A{cr0 - 1}"] = "Serie per il grafico (ogni 200 pezzi)"
J[f"A{cr0 - 1}"].font = BOLD
header(J, cr0, ["Pezzi", "Probabilità", ""])
for i, q in enumerate(range(0, 12001, 200)):
    r = cr0 + 1 + i
    cell(J, f"A{r}", q, fmt="#,##0")
    cell(J, f"B{r}",
         f"=IF(A{r}<{p('jp_min_qty')},0,MIN({p('jp_cap')},{p('jp_base')}+INT((A{r}-{p('jp_min_qty')})/{p('jp_step_qty')})*{p('jp_step')}))",
         fmt="0%")
last_c = cr0 + 1 + len(range(0, 12001, 200)) - 1
lc = LineChart()
lc.title = "Probabilità di jackpot e pezzi venduti"
lc.y_axis.title = "Probabilità"
lc.x_axis.title = "Pezzi venduti in un «Vendi tutto»"
lc.y_axis.number_format = "0%"
lc.y_axis.scaling.min = 0
lc.y_axis.scaling.max = 0.6
lc.add_data(Reference(J, min_col=2, min_row=cr0, max_row=last_c), titles_from_data=True)
lc.set_categories(Reference(J, min_col=1, min_row=cr0 + 1, max_row=last_c))
lc.series[0].graphicalProperties.line.solidFill = "EB6834"
lc.series[0].graphicalProperties.line.width = 25000
lc.legend = None
lc.height, lc.width = 9, 18
J.add_chart(lc, "E4")

# ---------------------------------------------------------------- Pity
T = sheet("Pity", (12, 14, 14, 16, 14, 16, 14, 50))
title(T, "Pity", "Ogni tentativo a vuoto con almeno 3.000 pezzi aggiunge +25% al tentativo successivo; dopo 3 a vuoto di fila il 4° è garantito. Si azzera al jackpot.")
T["A3"] = ("Nelle colonne blu scegli i pezzi e l'esito ipotetico di ogni tentativo (sì = jackpot): le probabilità e il contatore si ricalcolano. "
           "Un tentativo sotto i 3.000 pezzi non tocca il contatore (vedi Punti aperti).")
T["A3"].font = NOTE
scenarios = [
    ("A — sfortuna piena: 3 tentativi grossi a vuoto, il 4° garantito", [(5000, "no"), (5000, "no"), (5000, "no"), (5000, "sì")]),
    ("B — jackpot al secondo tentativo: il contatore si azzera", [(3000, "no"), (3000, "sì"), (3000, "no"), (3000, "no")]),
    ("C — tentativi piccoli non contano per il pity", [(1000, "no"), (3500, "no"), (800, "no"), (3500, "no"), (3500, "no"), (3500, "sì")]),
    ("D — vendita enorme: 50% + pity", [(10000, "no"), (10000, "no"), (10000, "sì")]),
    ("E — fortuna: jackpot al primo tentativo", [(400, "sì"), (400, "no")]),
]
r = 5
for name, steps in scenarios:
    T[f"A{r}"] = name
    T[f"A{r}"].font = BOLD
    r += 1
    header(T, r, ["Tentativo", "Pezzi", "Prob. base", "Conta per pity?", "Bonus pity", "Prob. totale", "Jackpot?", "Contatore a vuoto dopo"])
    first = r + 1
    for k, (q, hit) in enumerate(steps):
        rr = first + k
        cell(T, f"A{rr}", k + 1)
        cell(T, f"B{rr}", q, BLUE, "#,##0")
        cell(T, f"C{rr}",
             f"=IF(B{rr}<{p('jp_min_qty')},0,MIN({p('jp_cap')},{p('jp_base')}+INT((B{rr}-{p('jp_min_qty')})/{p('jp_step_qty')})*{p('jp_step')}))",
             fmt="0%")
        cell(T, f"D{rr}", f'=IF(B{rr}>={p("pity_qty")},"sì","no")')
        prev = f"H{rr - 1}" if k > 0 else "0"
        cell(T, f"E{rr}", f"={prev}*{p('pity_add')}", fmt="0%")
        cell(T, f"F{rr}", f"=IF({prev}>={p('pity_guarantee')},1,MIN(1,C{rr}+E{rr}))", fmt="0%")
        cell(T, f"G{rr}", hit, BLUE)
        cell(T, f"H{rr}", f'=IF(OR(G{rr}="sì",F{rr}>=1),0,IF(D{rr}="sì",{prev}+1,{prev}))')
    r = first + len(steps) + 2
T[f"A{r}"] = "Colonna «Jackpot?» con prob. totale 100%: il jackpot è certo anche se scrivi «no» (il contatore lo azzera comunque)."
T[f"A{r}"].font = NOTE

# ---------------------------------------------------------------- Esempi di vendita
E = sheet("Esempi di vendita", (44, 12, 12, 12, 12, 12, 12, 14, 14, 12, 16, 16))
title(E, "Esempi di vendita", "Prezzo pagato = prezzo del momento del clic per tutti i pezzi (decisione A). Lastra di ferro se non indicato. Le colonne blu si possono cambiare.")
header(E, 4, ["Caso", "B ($)", "Prezzo al clic (× B)", "Pezzi", "Schermo?", "«Vendi tutto»?", "Nel picco?",
              "Prezzo dopo (× B)", "Crollo?", "LED", "Jackpot possibile?", "Incasso ($)"])
b_plate = "'Valori B'!$F$6"  # iron-plate row (market order: iron-ore, iron-plate, ...)
cases = [
    ("1. Vendita piccola nella fascia normale", 1.3, 50, "no", "no"),
    ("2. 300 pezzi a prezzo normale: arriva a B e crolla", 1.3, 300, "no", "sì"),
    ("3. 300 pezzi nel picco: niente crollo", 2.9, 300, "sì", "sì"),
    ("4. 2.000 pezzi nel picco con lo schermo, jackpot mancato", 2.9, 2000, "sì", "sì"),
    ("5. 2.000 pezzi fuori dal picco: crollo, jackpot impossibile", 1.4, 2000, "sì", "sì"),
    ("6. 5.000 pezzi nel picco senza schermo: niente jackpot", 2.9, 5000, "no", "sì"),
    ("7. Vendita automatica di 100 pezzi", None, 100, "no", "no"),
]
for i, (name, price, q, screen, all_) in enumerate(cases):
    r = 5 + i
    cell(E, f"A{r}", name)
    cell(E, f"B{r}", f"={b_plate}", GREEN, "$#,##0")
    if price is None:
        cell(E, f"C{r}", f"={p('auto_price')}", GREEN, "0.00")
    else:
        cell(E, f"C{r}", price, BLUE, "0.00")
    cell(E, f"D{r}", q, BLUE, "#,##0")
    cell(E, f"E{r}", screen, BLUE)
    cell(E, f"F{r}", all_, BLUE)
    cell(E, f"G{r}", f'=IF(C{r}>={p("spike_in")},"sì","no")')
    if price is None:
        cell(E, f"H{r}", f"=C{r}", fmt="0.00")
        cell(E, f"I{r}", "no")
        cell(E, f"J{r}", "—")
        cell(E, f"K{r}", "no")
        cell(E, f"L{r}", f"=D{r}*ROUND(B{r}*C{r},0)", fmt="$#,##0")
        continue
    cell(E, f"I{r}", f'=IF(OR(D{r}>{p("crash_qty")},C{r}-D{r}*{p("impact")}<={p("norm_min")}),"sì","no")')
    cell(E, f"H{r}", f'=IF(I{r}="sì",{p("norm_min")},C{r}-D{r}*{p("impact")})', fmt="0.00",
         comment="Con il crollo il prezzo va vicino a B (e nei primi 5 minuti a volte sotto)." if i == 1 else None)
    cell(E, f"J{r}",
         f'=IF(AND(E{r}="sì",F{r}="sì",D{r}>={p("jp_min_qty")}),IF(I{r}="sì","giallo",IF(C{r}-H{r}<={p("led_mild")},"blu","viola")),"—")')
    cell(E, f"K{r}", f'=IF(AND(E{r}="sì",F{r}="sì",G{r}="sì",D{r}>={p("jp_min_qty")}),"sì","no")')
    cell(E, f"L{r}", f"=D{r}*ROUND(B{r}*C{r},0)", fmt="$#,##0")
last_e = 4 + len(cases)
jr = last_e + 3
E[f"A{jr - 1}"] = "Jackpot: stesso caso 4, ma esce il jackpot"
E[f"A{jr - 1}"].font = BOLD
header(E, jr, ["Voce", "Per pezzo ($)", "Pezzi", "Totale ($)", "", "", "", "", "", "", "", ""])
cell(E, f"A{jr + 1}", "Valore base B")
cell(E, f"B{jr + 1}", f"=B8", fmt="$#,##0")
cell(E, f"A{jr + 2}", "Bonus +777% di B (per difetto)")
cell(E, f"B{jr + 2}", f"=ROUNDDOWN(B8*{p('jp_bonus')},0)", fmt="$#,##0")
cell(E, f"A{jr + 3}", "Rialzo del momento (prezzo − B)")
cell(E, f"B{jr + 3}", f"=ROUND(B8*C8,0)-B8", fmt="$#,##0")
cell(E, f"A{jr + 4}", "Totale jackpot", BOLD)
cell(E, f"B{jr + 4}", f"=SUM(B{jr + 1}:B{jr + 3})", BOLD, "$#,##0")
for k in range(1, 5):
    cell(E, f"C{jr + k}", "=D8", fmt="#,##0")
    cell(E, f"D{jr + k}", f"=B{jr + k}*C{jr + k}", BOLD if k == 4 else BLACK, "$#,##0")
cell(E, f"A{jr + 5}", "Senza jackpot (caso 4) per confronto")
cell(E, f"D{jr + 5}", "=L8", fmt="$#,##0")
cell(E, f"A{jr + 6}", "Rapporto jackpot / vendita normale nel picco")
cell(E, f"D{jr + 6}", f"=D{jr + 4}/D{jr + 5}", fmt="0.0x")
ex = jr + 8
E[f"A{ex}"] = "Esempio dell'utente: B 10, prezzo 30 (+200%) → 10 + 77 + 20 ="
E[f"A{ex}"].font = NOTE
cell(E, f"B{ex}", f"=10+ROUNDDOWN(10*{p('jp_bonus')},0)+(30-10)", BOLD, "$#,##0")
leg = ex + 2
for k, line in enumerate([
    "LED dello schermo (sequenza solo per «Vendi tutto» da 400 pezzi in su, con lo schermo):",
    "blu = calo lieve (fino al 10% di B) · viola = calo grave senza crollo · giallo = crollo (jackpot possibile se nel picco).",
    "Attenzione: con la soglia del crollo a 400 pezzi, ogni vendita che avvia la sequenza fa crollare → il LED è sempre giallo. Vedi Punti aperti.",
]):
    E[f"A{leg + k}"] = line
    E[f"A{leg + k}"].font = NOTE if k else BOLD

# ---------------------------------------------------------------- Simulazione 1 ora
S = sheet("Simulazione 1 ora", (12, 14, 14, 30))
title(S, "Un'ora di Lastra di ferro (simulazione illustrativa)",
      "Punti ogni 10 s generati da uno script con seme fisso: andamento normale, un picco al minuto 14, un «Vendi tutto» di 2.000 pezzi fuori dal picco al minuto 33 → crollo di 20 minuti.")
header(S, 4, ["Minuto", "Prezzo (× B)", "Prezzo ($)", "Fase"])
rng = random.Random(268)
norm_lo, norm_hi, avg = 1.0, 1.6, 1.3
level = 1.3
rows = []
crash_start = 33 * 60
spike_start = 14 * 60
for step in range(0, 3601, 10):
    t = step
    phase = "normale"
    # Mean-reverting wander inside the normal band.
    level += (avg - level) * 0.08 + rng.uniform(-0.06, 0.06)
    level = min(norm_hi, max(norm_lo, level))
    price = level
    if spike_start <= t <= spike_start + 50:
        s = t - spike_start
        top = 2.93
        if s <= 10:
            price = level + (top - level) * s / 10
        elif s <= 20:
            price = top
        else:
            price = top - (top - level) * (s - 20) / 30
        phase = "picco"
    if t == crash_start:
        price = level - 2000 * 0.001
        phase = "vendita 2.000 pezzi"
    if crash_start < t <= crash_start + 20 * 60:
        phase = "crollo"
        price = 1.0 + rng.uniform(0, 0.04)
        if t <= crash_start + 5 * 60 and rng.random() < 0.25:
            price = rng.uniform(0.6, 0.7)
            phase = "crollo, sotto B"
        level = 1.0
    if t > crash_start + 20 * 60 and phase == "normale":
        level = min(level, 1.15) if t < crash_start + 21 * 60 else level
        price = level
    rows.append((t / 60, round(max(0.6, price), 3), phase))
for i, (m, price, phase) in enumerate(rows):
    r = 5 + i
    cell(S, f"A{r}", m, fmt="0.0")
    cell(S, f"B{r}", price, BLUE, "0.00")
    cell(S, f"C{r}", f"=ROUND(B{r}*{b_plate},0)", fmt="$#,##0")
    cell(S, f"D{r}", phase)
S["B5"].comment = Comment("Valori generati da uno script Python con seme fisso (268): illustrano le regole, non sono il codice del gioco.", "Claude")
last_s = 4 + len(rows)
lc = LineChart()
lc.title = "Lastra di ferro: un'ora di mercato"
lc.y_axis.title = "Prezzo ($)"
lc.x_axis.title = "Minuto"
lc.add_data(Reference(S, min_col=3, min_row=4, max_row=last_s), titles_from_data=True)
lc.set_categories(Reference(S, min_col=1, min_row=5, max_row=last_s))
lc.series[0].graphicalProperties.line.solidFill = "2A78D6"
lc.series[0].graphicalProperties.line.width = 19050
lc.series[0].smooth = False
lc.x_axis.tickLblSkip = 30
lc.x_axis.number_format = "0"
lc.legend = None
lc.height, lc.width = 10, 24
S.add_chart(lc, "F4")

# ---------------------------------------------------------------- Ramo Finanza
R = sheet("Ramo Finanza", (30, 16, 40, 22, 20))
title(R, "Ramo Finanza della Ricerca (specifica 2)",
      "Moduli da incollare allo schermo di borsa. Costi e prerequisiti esatti: da decidere (sfondo giallo).")
header(R, 4, ["Blocco / modulo", "Fase di gioco", "Cosa fa", "Prerequisito", "Costo"])
mods = [
    ("Schermo di borsa + bot della finanza", "inizio Finanza", "Rettangolo senza gambe accanto al CORE. Si imposta su un oggetto da monitorare (grafico). Senza schermo il jackpot non esiste. «Vendi tutto» da 400 pezzi avvia la sequenza del drone 2.5D.", "Mercato (CORE)", "da decidere"),
    ("Modulo previsione (drone)", "primi livelli", "Il drone prevede l'andamento e consiglia quando vendere. Non prevede i jackpot.", "Schermo di borsa", "da decidere"),
    ("Modulo percentuale", "metà partita", "Mostra la percentuale attuale rispetto al valore base B.", "Modulo previsione", "da decidere"),
    ("Modulo probabilità jackpot", "fine partita", "Mostra la probabilità attuale di jackpot facendo «Vendi tutto» (pity compreso).", "Modulo percentuale", "da decidere"),
]
for i, row in enumerate(mods):
    r = 5 + i
    for j, v in enumerate(row):
        c = cell(R, f"{get_column_letter(j + 1)}{r}", v, fill=YELLOW if j == 4 else None)
        c.alignment = WRAP
    R.row_dimensions[r].height = 60
seq = 11
R[f"A{seq - 1}"] = "Sequenza del «Vendi tutto» con lo schermo (decisa dall'utente; dettagli della tensione da discutere)"
R[f"A{seq - 1}"].font = BOLD
header(R, seq, ["Passo", "Durata", "Cosa si vede", "", ""])
steps = [
    ("1. Zoom senza HUD", "—", "La visuale va sul drone in 2.5D davanti allo schermo; l'HUD sparisce."),
    ("2. Cambio canale", "—", "Il drone cambia canale sullo schermo ed è sorpreso."),
    ("3. Corsa al CORE", "—", "Corre al CORE e riempie la valigetta con la quantità venduta."),
    ("4. Ritorno", "—", "Torna di fretta. Il grafico va avanti al rallentatore, fermo sul prezzo del clic."),
    ("5. Lancio e «calcolo…»", "2 s", "Lancia la valigetta; sullo schermo «calcolo…»."),
    ("6. LED", "2 s", "La striscia gialla e nera sotto lo schermo si accende: blu (lieve), viola (grave), giallo (crollo)."),
    ("7. Risultato", "—", "Incasso sullo schermo. Jackpot: luci gialle rotanti negli angoli della cornice, coriandoli, schermata JACKPOT."),
]
for i, (a, b, c) in enumerate(steps):
    r = seq + 1 + i
    cell(R, f"A{r}", a)
    cell(R, f"B{r}", b)
    cell(R, f"C{r}", c).alignment = WRAP
    R.row_dimensions[r].height = 32
R[f"A{seq + 9}"] = "Il risultato si decide al clic (simulazione deterministica); l'animazione lo mostra soltanto, la fabbrica continua a lavorare."
R[f"A{seq + 9}"].font = NOTE

# ---------------------------------------------------------------- Magazzino CORE
M = sheet("Magazzino CORE", (12, 22, 22, 46))
title(M, "Magazzino del CORE e drone per livello (specifica 3, anteprima)",
      "Massimo per oggetto: quando un oggetto è al massimo non entra più e le linee che lo portano si fermano, come in Mindustry. Numeri proposti, da decidere.")
header(M, 4, ["Livello CORE", "Massimo per oggetto", "Drone", "Cosa permette al Mercato"])
levels = [(1, 1000, "drone-1"), (2, 4000, "drone-2 (più veloce)"), (3, 12000, "drone-3 (ancora più veloce)")]
for i, (lvl, cap, drone) in enumerate(levels):
    r = 5 + i
    cell(M, f"A{r}", lvl)
    cell(M, f"B{r}", cap, BLUE, "#,##0", YELLOW)
    cell(M, f"C{r}", drone)
    cell(M, f"D{r}",
         f'="Vendi tutto fino a "&TEXT(B{r},"#,##0")&" pezzi: jackpot fino al "&TEXT(IF(B{r}<{p("jp_min_qty")},0,MIN({p("jp_cap")},{p("jp_base")}+INT((B{r}-{p("jp_min_qty")})/{p("jp_step_qty")})*{p("jp_step")})),"0%")&IF(B{r}>={p("pity_qty")},", pity attivo",", pity non raggiungibile")').alignment = WRAP
M["A9"] = "Il magazzino fa da freno naturale: per il pity servono 3.000 pezzi e per il 50% ne servono 10.000, quindi i jackpot grossi arrivano solo con un CORE potenziato."
M["A9"].font = NOTE

# ---------------------------------------------------------------- Punti aperti
O = sheet("Punti aperti", (6, 100))
title(O, "Punti aperti", "Da decidere prima della specifica (o nelle micro correzioni dopo).")
openq = [
    "Incoerenza: la sequenza parte solo con «Vendi tutto» da 400 pezzi, ma oltre 400 pezzi c'è sempre il crollo: i LED sarebbero sempre gialli e blu/viola non comparirebbero mai. Proposte: (a) la sequenza parte con ogni «Vendi tutto» quando c'è lo schermo; (b) la soglia del crollo sale (es. 1.000 pezzi) e resta 400 quella del jackpot.",
    "Pity: un tentativo sotto i 3.000 pezzi azzera la serie «di fila» o la lascia com'è? Nel foglio Pity la lascia com'è.",
    "Pity: conta per oggetto (come il resto del mercato) o per partita? Nel foglio è per oggetto.",
    "Pity: il bonus accumulato vale anche per un tentativo sotto i 3.000 pezzi (nel foglio sì, scenario C tentativo 3) o solo per quelli da 3.000 in su?",
    "Il jackpot provoca il crollo di 20 minuti come una vendita normale? (Detto di sì: «il solito crollo».)",
    "Valori B: moltiplicatore e lista degli oggetti di fine partita (foglio Valori B), da rifare insieme alla campagna.",
    "Soglie dei LED: blu fino al 10% di calo di B, oltre viola.",
    "Impatto delle vendite piccole (0,1% di B per pezzo) e recupero in 2 minuti: valori di partenza da provare in partita.",
    "Magazzino del CORE: massimi per livello (1.000 / 4.000 / 12.000) e quanti livelli.",
    "Ramo Finanza: costi e prerequisiti dei moduli; eventuali potenziamenti dello schermo oltre i tre moduli.",
    "Agenti (server MCP): devono vedere prezzi, fasi e picchi come la panoramica, e non i moduli che il giocatore non ha sbloccato?",
]
for i, q in enumerate(openq):
    r = 4 + i
    cell(O, f"A{r}", i + 1)
    c = cell(O, f"B{r}", q, fill=YELLOW if i == 0 else None)
    c.alignment = WRAP
    O.row_dimensions[r].height = 46 if len(q) > 110 else 20

for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.font is not None and c.font.name != FONT:
                c.font = Font(name=FONT, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)

for w in wb.worksheets:
    w.page_setup.orientation = "landscape"
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.page_setup.fitToWidth = 1
    w.page_setup.fitToHeight = 0

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("scritto", OUT)
